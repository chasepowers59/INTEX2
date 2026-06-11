#!/usr/bin/env python3
"""Build an editable 2026 card-spending workbook from uploaded PDF statements."""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from pypdf import PdfReader


UPLOAD_DIR = Path(
    os.environ.get("CARD_PDF_DIR", "/home/ubuntu/.cursor/projects/workspace/uploads")
)
OUTPUT_PATH = Path(
    os.environ.get(
        "CARD_WORKBOOK_PATH",
        "/workspace/output/card-spending/2026_card_spending_by_month.xlsx",
    )
)

CITI_FILES = [
    "January2026_f0b0.pdf",
    "February2026_f8d6.pdf",
    "March2026_da89.pdf",
    "April2026_4952.pdf",
    "May2026_2c5d.pdf",
    "June2026_4656.pdf",
]

CHASE_FILES = [
    "F25604AB-F6FC-4AC2-96CD-5E3363840815-list_adc9.pdf",
    "B3CDA793-1958-4FE0-A464-444F7501BD7F-list_3284.pdf",
    "36C0F7B9-D67A-4AE6-9684-B5F848CBA9D5-list_4699.pdf",
    "BE3F4263-FAF2-4A9C-87A5-53852C30604D-list_09af.pdf",
    "69212C10-1B95-4A05-9104-379BBA52768F-list_d242.pdf",
]

MONTH_SHEETS = [f"{month:02d}월" for month in range(1, 13)]
DATA_HEADERS = [
    "ID",
    "사용일",
    "전기일",
    "카드",
    "거래유형",
    "상점/원본설명",
    "표준상점",
    "원본_대분류",
    "원본_세부분류",
    "상세태그",
    "원본금액",
    "지출금액",
    "환불/크레딧",
    "결제금액",
    "순액",
    "명세서월",
    "원본파일",
    "사용자_대분류",
    "사용자_세부분류",
    "사용자_메모",
    "검토필요",
]

CATEGORIES = [
    "식료품/생활용품",
    "교통/주유",
    "온라인쇼핑",
    "보험",
    "엔터테인먼트",
    "결제/이체",
    "환불/크레딧",
    "기타/확인필요",
]

SUBCATEGORIES = [
    "Costco 매장",
    "Costco 온라인",
    "주유",
    "Amazon",
    "자동차보험",
    "게임/디지털",
    "카드대금 결제",
    "환불/반품",
    "기타",
]

TRANSACTION_TYPES = ["구매", "환불/크레딧", "결제"]
DATE_RE = re.compile(r"^\d{2}/\d{2}$")
MONEY_TOKEN_RE = re.compile(r"^[+-]?\$?(?:\d[\d,]*|\d*\.\d+|\.\d+)(?:\.\d+)?$")
CHASE_TX_RE = re.compile(r"^(\d{2}/\d{2})\s+(.+?)\s+([+-]?(?:\d[\d,]*|\.\d+|\d*\.\d+))$")


@dataclass
class Transaction:
    transaction_date: date
    post_date: date | None
    card: str
    transaction_type: str
    description: str
    merchant: str
    category: str
    subcategory: str
    tag: str
    amount: Decimal
    statement_month: str
    source_file: str
    needs_review: str


@dataclass
class StatementSummary:
    source_file: str
    card: str
    statement_month: str
    period_start: date
    period_end: date
    statement_purchases: Decimal | None
    parsed_purchases: Decimal
    parsed_credits: Decimal
    parsed_payments: Decimal
    transaction_count: int


def money(value: str) -> Decimal:
    cleaned = value.strip().replace("$", "").replace(",", "")
    if cleaned.startswith("+"):
        cleaned = cleaned[1:]
    if cleaned.startswith("-."):
        cleaned = "-0" + cleaned[1:]
    elif cleaned.startswith("."):
        cleaned = "0" + cleaned
    return Decimal(cleaned).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def is_money_line(value: str) -> bool:
    value = value.strip()
    if value in {"-", "+"}:
        return False
    try:
        money(value)
        return bool(MONEY_TOKEN_RE.match(value.replace("$", "")) or "$" in value)
    except Exception:
        return False


def read_pdf_lines(path: Path) -> list[str]:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return [line.strip() for line in text.splitlines() if line.strip()]


def parse_mmdd_year(mmdd: str, period_start: date, period_end: date) -> date:
    month, day = [int(part) for part in mmdd.split("/")]
    if period_start.year != period_end.year:
        year = period_start.year if month >= period_start.month else period_end.year
    else:
        year = period_end.year
    return date(year, month, day)


def parse_short_date(value: str) -> date:
    return datetime.strptime(value, "%m/%d/%y").date()


def infer_statement_month(period_end: date) -> str:
    return f"{period_end.year}-{period_end.month:02d}"


def normalize_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def categorize(description: str, transaction_type: str) -> tuple[str, str, str, str, str]:
    upper = description.upper()
    if transaction_type == "결제":
        return ("Card Payment", "결제/이체", "카드대금 결제", "카드결제", "")
    if transaction_type == "환불/크레딧":
        if "COSTCO" in upper:
            return ("Costco", "환불/크레딧", "환불/반품", "Costco 환불", "")
        if "AMAZON" in upper:
            return ("Amazon", "환불/크레딧", "환불/반품", "Amazon 환불", "")
        return ("Refund/Credit", "환불/크레딧", "환불/반품", "환불/크레딧", "")
    if "COSTCO GAS" in upper:
        return ("Costco Gas", "교통/주유", "주유", "Costco Gas", "")
    if "COSTCO WHSE" in upper:
        return ("Costco Warehouse", "식료품/생활용품", "Costco 매장", "Costco 매장", "")
    if "WWW COSTCO" in upper or "COSTCO COM" in upper:
        return ("Costco Online", "온라인쇼핑", "Costco 온라인", "Costco 온라인", "")
    if "AMAZON" in upper or "AMZN.COM" in upper:
        return ("Amazon", "온라인쇼핑", "Amazon", "Amazon", "")
    if "PROGRESSIVE INS" in upper:
        return ("Progressive Insurance", "보험", "자동차보험", "자동차보험", "")
    if "NINTENDO" in upper:
        return ("Nintendo", "엔터테인먼트", "게임/디지털", "게임/디지털", "")
    return ("기타", "기타/확인필요", "기타", "검토필요", "확인")


def make_transaction(
    *,
    transaction_date: date,
    post_date: date | None,
    card: str,
    transaction_type: str,
    description: str,
    amount: Decimal,
    statement_month: str,
    source_file: str,
) -> Transaction:
    merchant, category, subcategory, tag, review = categorize(description, transaction_type)
    return Transaction(
        transaction_date=transaction_date,
        post_date=post_date,
        card=card,
        transaction_type=transaction_type,
        description=normalize_spaces(description),
        merchant=merchant,
        category=category,
        subcategory=subcategory,
        tag=tag,
        amount=amount,
        statement_month=statement_month,
        source_file=source_file,
        needs_review=review,
    )


def find_first_index(lines: list[str], predicates: Iterable[str]) -> int | None:
    predicate_tuple = tuple(predicates)
    for index, line in enumerate(lines):
        if line in predicate_tuple or any(line.startswith(prefix) for prefix in predicate_tuple):
            return index
    return None


def parse_citi_statement(path: Path) -> tuple[list[Transaction], StatementSummary]:
    lines = read_pdf_lines(path)
    text = "\n".join(lines)
    period_match = re.search(r"Billing Period:\s*(\d{2}/\d{2}/\d{2})-(\d{2}/\d{2}/\d{2})", text)
    if not period_match:
        raise ValueError(f"Could not find Citi billing period in {path.name}")
    period_start = parse_short_date(period_match.group(1))
    period_end = parse_short_date(period_match.group(2))
    statement_month = infer_statement_month(period_end)
    card = "Costco Anywhere Visa (Citi) ****6858"

    statement_purchases = None
    purchases_match = re.search(r"Purchases\s+\+?(\$[\d,]+\.\d{2})", text)
    if purchases_match:
        statement_purchases = money(purchases_match.group(1))

    transactions: list[Transaction] = []
    pay_idx = find_first_index(lines, ["Payments, Credits and Adjustments"])
    purchase_idx = find_first_index(lines, ["Standard Purchases", "Purchases Prior", "Purchases After"])
    if pay_idx is None or purchase_idx is None:
        return transactions, StatementSummary(
            path.name,
            card,
            statement_month,
            period_start,
            period_end,
            statement_purchases,
            Decimal("0.00"),
            Decimal("0.00"),
            Decimal("0.00"),
            0,
        )

    i = pay_idx + 1
    while i < purchase_idx:
        if DATE_RE.match(lines[i]) and i + 2 < purchase_idx:
            tx_date = parse_mmdd_year(lines[i], period_start, period_end)
            description = lines[i + 1]
            sign = Decimal("1.00")
            amount_index = i + 2
            if lines[i + 2] == "-":
                sign = Decimal("-1.00")
                amount_index = i + 3
            if amount_index < purchase_idx and is_money_line(lines[amount_index]):
                amount = money(lines[amount_index]) * sign
                tx_type = "결제" if "PAYMENT" in description.upper() or "AUTOPAY" in description.upper() else "환불/크레딧"
                transactions.append(
                    make_transaction(
                        transaction_date=tx_date,
                        post_date=None,
                        card=card,
                        transaction_type=tx_type,
                        description=description,
                        amount=amount,
                        statement_month=statement_month,
                        source_file=path.name,
                    )
                )
                i = amount_index + 1
                continue
        i += 1

    i = purchase_idx + 1
    while i < len(lines):
        if lines[i] in {"Fees Charged", "Interest Charged"}:
            break
        if DATE_RE.match(lines[i]) and i + 3 < len(lines) and DATE_RE.match(lines[i + 1]) and is_money_line(lines[i + 3]):
            tx_date = parse_mmdd_year(lines[i], period_start, period_end)
            post_date = parse_mmdd_year(lines[i + 1], period_start, period_end)
            description = lines[i + 2]
            amount = money(lines[i + 3])
            transactions.append(
                make_transaction(
                    transaction_date=tx_date,
                    post_date=post_date,
                    card=card,
                    transaction_type="구매",
                    description=description,
                    amount=amount,
                    statement_month=statement_month,
                    source_file=path.name,
                )
            )
            i += 4
            continue
        i += 1

    return transactions, build_summary(path.name, card, statement_month, period_start, period_end, statement_purchases, transactions)


def parse_chase_statement(path: Path) -> tuple[list[Transaction], StatementSummary]:
    lines = read_pdf_lines(path)
    text = "\n".join(lines)
    period_match = re.search(r"Opening/Closing Date\s+(\d{2}/\d{2}/\d{2})\s*-\s*(\d{2}/\d{2}/\d{2})", text)
    if not period_match:
        raise ValueError(f"Could not find Chase opening/closing date in {path.name}")
    period_start = parse_short_date(period_match.group(1))
    period_end = parse_short_date(period_match.group(2))
    statement_month = infer_statement_month(period_end)
    card = "Prime Visa (Chase) ****6217"

    statement_purchases = None
    purchases_match = re.search(r"Purchases\s+\+?(\$?[\d,]*\.?\d+\.\d{2}|\$?0\.00)", text)
    if purchases_match:
        statement_purchases = money(purchases_match.group(1))

    start_idx = None
    for index, line in enumerate(lines):
        if line.startswith("Transaction Merchant") and "$ Amount" in line:
            start_idx = index + 1
            break

    transactions: list[Transaction] = []
    if start_idx is None:
        return transactions, build_summary(path.name, card, statement_month, period_start, period_end, statement_purchases, transactions)

    for line in lines[start_idx:]:
        if line.startswith("Total fees charged") or line.startswith("Total interest charged") or line.startswith("Purchases "):
            break
        if line.startswith("Transaction Merchant") or line.startswith("Split Transaction"):
            break
        match = CHASE_TX_RE.match(line)
        if not match:
            continue
        tx_date = parse_mmdd_year(match.group(1), period_start, period_end)
        description = match.group(2)
        amount = money(match.group(3))
        upper_desc = description.upper()
        if "PAYMENT THANK YOU" in upper_desc:
            tx_type = "결제"
        elif amount < 0:
            tx_type = "환불/크레딧"
        else:
            tx_type = "구매"
        transactions.append(
            make_transaction(
                transaction_date=tx_date,
                post_date=None,
                card=card,
                transaction_type=tx_type,
                description=description,
                amount=amount,
                statement_month=statement_month,
                source_file=path.name,
            )
        )

    return transactions, build_summary(path.name, card, statement_month, period_start, period_end, statement_purchases, transactions)


def build_summary(
    source_file: str,
    card: str,
    statement_month: str,
    period_start: date,
    period_end: date,
    statement_purchases: Decimal | None,
    transactions: list[Transaction],
) -> StatementSummary:
    parsed_purchases = sum((tx.amount for tx in transactions if tx.transaction_type == "구매"), Decimal("0.00"))
    parsed_credits = sum((tx.amount for tx in transactions if tx.transaction_type == "환불/크레딧"), Decimal("0.00"))
    parsed_payments = sum((tx.amount for tx in transactions if tx.transaction_type == "결제"), Decimal("0.00"))
    return StatementSummary(
        source_file=source_file,
        card=card,
        statement_month=statement_month,
        period_start=period_start,
        period_end=period_end,
        statement_purchases=statement_purchases,
        parsed_purchases=parsed_purchases,
        parsed_credits=parsed_credits,
        parsed_payments=parsed_payments,
        transaction_count=len(transactions),
    )


def load_transactions() -> tuple[list[Transaction], list[StatementSummary]]:
    transactions: list[Transaction] = []
    summaries: list[StatementSummary] = []
    for file_name in CITI_FILES:
        parsed, summary = parse_citi_statement(UPLOAD_DIR / file_name)
        transactions.extend(parsed)
        summaries.append(summary)
    for file_name in CHASE_FILES:
        parsed, summary = parse_chase_statement(UPLOAD_DIR / file_name)
        transactions.extend(parsed)
        summaries.append(summary)
    transactions.sort(key=lambda tx: (tx.transaction_date, tx.card, tx.description, tx.amount))
    return transactions, summaries


def apply_common_sheet_format(ws, max_col: int) -> None:
    for column_cells in ws.iter_cols(min_col=1, max_col=max_col):
        letter = column_cells[0].column_letter
        ws.column_dimensions[letter].width = 14
    for cell in ws[10]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A11"


def write_transaction_sheet(
    wb: Workbook,
    title: str,
    transactions: list[Transaction],
    subtitle: str,
    make_table: bool = True,
) -> None:
    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "구매 지출 합계"
    ws["A4"] = "환불/크레딧 합계"
    ws["A5"] = "카드대금 결제 합계"
    ws["A6"] = "순액"
    ws["A7"] = "거래 수"
    ws["D3"] = "편집 안내"
    ws["E3"] = "사용자_대분류/사용자_세부분류/사용자_메모 열은 직접 수정용입니다. 원본 분류는 비교를 위해 보존했습니다."

    header_row = 10
    for col, header in enumerate(DATA_HEADERS, start=1):
        ws.cell(row=header_row, column=col, value=header)

    start_row = header_row + 1
    for offset, tx in enumerate(transactions):
        row = start_row + offset
        tx_id = f"{tx.transaction_date:%Y%m%d}-{row - header_row:04d}"
        values = [
            tx_id,
            tx.transaction_date,
            tx.post_date,
            tx.card,
            tx.transaction_type,
            tx.description,
            tx.merchant,
            tx.category,
            tx.subcategory,
            tx.tag,
            float(tx.amount),
            f'=IF($E{row}="구매",$K{row},0)',
            f'=IF($E{row}="환불/크레딧",$K{row},0)',
            f'=IF($E{row}="결제",$K{row},0)',
            f"=$K{row}",
            tx.statement_month,
            tx.source_file,
            "",
            "",
            "",
            tx.needs_review,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)

    last_row = max(header_row, start_row + len(transactions) - 1)
    ws["B3"] = f"=SUM(L{start_row}:L{last_row})" if transactions else 0
    ws["B4"] = f"=SUM(M{start_row}:M{last_row})" if transactions else 0
    ws["B5"] = f"=SUM(N{start_row}:N{last_row})" if transactions else 0
    ws["B6"] = f"=SUM(O{start_row}:O{last_row})" if transactions else 0
    ws["B7"] = len(transactions)

    ws.auto_filter.ref = f"A{header_row}:U{last_row}"
    if make_table and transactions:
        table = Table(displayName=re.sub(r"\W+", "", f"Table_{title}"), ref=f"A{header_row}:U{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)

    apply_common_sheet_format(ws, len(DATA_HEADERS))
    widths = {
        "A": 18,
        "B": 12,
        "C": 12,
        "D": 32,
        "E": 14,
        "F": 46,
        "G": 24,
        "H": 18,
        "I": 18,
        "J": 18,
        "K": 13,
        "L": 13,
        "M": 15,
        "N": 15,
        "O": 13,
        "P": 13,
        "Q": 42,
        "R": 18,
        "S": 18,
        "T": 32,
        "U": 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in range(start_row, last_row + 1):
        for col in ["B", "C"]:
            ws[f"{col}{row}"].number_format = "yyyy-mm-dd"
        for col in ["K", "L", "M", "N", "O"]:
            ws[f"{col}{row}"].number_format = '$#,##0.00;[Red]-$#,##0.00'
    for row in range(3, 7):
        ws[f"B{row}"].number_format = '$#,##0.00;[Red]-$#,##0.00'

    category_dv = DataValidation(type="list", formula1="'분류목록'!$A$2:$A$9", allow_blank=True)
    subcategory_dv = DataValidation(type="list", formula1="'분류목록'!$B$2:$B$10", allow_blank=True)
    type_dv = DataValidation(type="list", formula1='"구매,환불/크레딧,결제"', allow_blank=False)
    ws.add_data_validation(category_dv)
    ws.add_data_validation(subcategory_dv)
    ws.add_data_validation(type_dv)
    if last_row >= start_row:
        category_dv.add(f"R{start_row}:R{last_row}")
        subcategory_dv.add(f"S{start_row}:S{last_row}")
        type_dv.add(f"E{start_row}:E{last_row}")
        ws.conditional_formatting.add(
            f"U{start_row}:U{last_row}",
            CellIsRule(operator="equal", formula=['"확인"'], fill=PatternFill("solid", fgColor="FFF2CC")),
        )


def write_category_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("분류목록")
    ws["A1"] = "대분류"
    ws["B1"] = "세부분류"
    ws["D1"] = "분류 기준"
    ws["D2"] = "Costco Gas -> 교통/주유 / 주유"
    ws["D3"] = "Costco WHSE -> 식료품/생활용품 / Costco 매장"
    ws["D4"] = "WWW COSTCO COM -> 온라인쇼핑 / Costco 온라인"
    ws["D5"] = "Amazon -> 온라인쇼핑 / Amazon"
    ws["D6"] = "Progressive Insurance -> 보험 / 자동차보험"
    ws["D7"] = "Nintendo -> 엔터테인먼트 / 게임/디지털"
    ws["D8"] = "결제 및 환불은 지출 합계와 별도로 집계"
    for row, category in enumerate(CATEGORIES, start=2):
        ws.cell(row=row, column=1, value=category)
    for row, subcategory in enumerate(SUBCATEGORIES, start=2):
        ws.cell(row=row, column=2, value=subcategory)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["D"].width = 52


def write_summary_sheet(wb: Workbook, transactions_2026: list[Transaction]) -> None:
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "2026년 카드 사용기록 요약"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = "월별 탭은 실제 사용일 기준이며, 2025년 12월 이월 거래는 별도 탭에 분리했습니다."
    headers = ["월", "구매 지출 합계", "환불/크레딧 합계", "카드대금 결제 합계", "순액", "거래 수"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=header)
    for month in range(1, 13):
        row = month + 4
        sheet = f"{month:02d}월"
        ws.cell(row=row, column=1, value=sheet)
        ws.cell(row=row, column=2, value=f"='{sheet}'!$B$3")
        ws.cell(row=row, column=3, value=f"='{sheet}'!$B$4")
        ws.cell(row=row, column=4, value=f"='{sheet}'!$B$5")
        ws.cell(row=row, column=5, value=f"='{sheet}'!$B$6")
        ws.cell(row=row, column=6, value=f"='{sheet}'!$B$7")

    ws["A19"] = "기본 대분류별 구매 지출"
    ws["A19"].font = Font(bold=True, size=13)
    ws["A20"] = "대분류"
    ws["B20"] = "구매 지출 합계"
    category_totals: dict[str, Decimal] = {}
    for tx in transactions_2026:
        if tx.transaction_type == "구매":
            category_totals[tx.category] = category_totals.get(tx.category, Decimal("0.00")) + tx.amount
    for row, (category, total) in enumerate(sorted(category_totals.items()), start=21):
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=2, value=float(total))

    ws["D19"] = "사용 방법"
    ws["D20"] = "1) 각 월 탭에서 필터로 상점/분류를 확인하세요."
    ws["D21"] = "2) 직접 바꾸고 싶은 분류는 사용자_대분류/사용자_세부분류 열에 입력하세요."
    ws["D22"] = "3) 원본_대분류/원본_세부분류는 자동 분류 결과라 비교용으로 남겨두었습니다."
    ws["D23"] = "4) 원본금액이나 거래유형을 바꾸면 월 탭 상단의 합계가 엑셀에서 다시 계산됩니다."

    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    for cell in ws[20]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    for column in ["B", "C", "D", "E"]:
        for row in range(5, 33):
            ws[f"{column}{row}"].number_format = '$#,##0.00;[Red]-$#,##0.00'
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 12


def write_statement_summary(wb: Workbook, summaries: list[StatementSummary]) -> None:
    ws = wb.create_sheet("원본_명세서요약")
    headers = [
        "원본파일",
        "카드",
        "명세서월",
        "시작일",
        "종료일",
        "명세서 구매합계",
        "추출 구매합계",
        "차이",
        "추출 환불/크레딧",
        "추출 결제",
        "거래 수",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    for row, summary in enumerate(summaries, start=2):
        statement_purchases = float(summary.statement_purchases) if summary.statement_purchases is not None else None
        ws.cell(row=row, column=1, value=summary.source_file)
        ws.cell(row=row, column=2, value=summary.card)
        ws.cell(row=row, column=3, value=summary.statement_month)
        ws.cell(row=row, column=4, value=summary.period_start)
        ws.cell(row=row, column=5, value=summary.period_end)
        ws.cell(row=row, column=6, value=statement_purchases)
        ws.cell(row=row, column=7, value=float(summary.parsed_purchases))
        ws.cell(row=row, column=8, value=f"=G{row}-F{row}" if statement_purchases is not None else "")
        ws.cell(row=row, column=9, value=float(summary.parsed_credits))
        ws.cell(row=row, column=10, value=float(summary.parsed_payments))
        ws.cell(row=row, column=11, value=summary.transaction_count)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    ws.auto_filter.ref = f"A1:K{max(1, len(summaries) + 1)}"
    for col in ["D", "E"]:
        for row in range(2, len(summaries) + 2):
            ws[f"{col}{row}"].number_format = "yyyy-mm-dd"
    for col in ["F", "G", "H", "I", "J"]:
        for row in range(2, len(summaries) + 2):
            ws[f"{col}{row}"].number_format = '$#,##0.00;[Red]-$#,##0.00'
    widths = {"A": 42, "B": 32, "C": 12, "D": 12, "E": 12, "F": 16, "G": 16, "H": 12, "I": 16, "J": 14, "K": 10}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def style_workbook(wb: Workbook) -> None:
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border


def build_workbook() -> tuple[Path, list[Transaction], list[StatementSummary]]:
    transactions, summaries = load_transactions()
    transactions_2026 = [tx for tx in transactions if tx.transaction_date.year == 2026]
    carryover = [tx for tx in transactions if tx.transaction_date.year != 2026]

    wb = Workbook()
    write_summary_sheet(wb, transactions_2026)
    write_category_sheet(wb)
    write_transaction_sheet(wb, "전체거래", transactions_2026, "2026년 전체 카드 거래 - 실제 사용일 기준")
    for month, sheet_name in enumerate(MONTH_SHEETS, start=1):
        monthly_transactions = [tx for tx in transactions_2026 if tx.transaction_date.month == month]
        write_transaction_sheet(wb, sheet_name, monthly_transactions, f"2026년 {month}월 카드 사용내역")
    write_transaction_sheet(wb, "2025이월거래", carryover, "2025년 사용일 거래 - 2026년 1월 명세서에 포함")
    write_statement_summary(wb, summaries)
    style_workbook(wb)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH, transactions, summaries


def main() -> None:
    output_path, transactions, summaries = build_workbook()
    print(f"Workbook: {output_path}")
    print(f"Extracted transactions: {len(transactions)}")
    print(f"2026 transactions: {sum(1 for tx in transactions if tx.transaction_date.year == 2026)}")
    print(f"Carryover transactions: {sum(1 for tx in transactions if tx.transaction_date.year != 2026)}")
    for summary in summaries:
        diff = None
        if summary.statement_purchases is not None:
            diff = summary.parsed_purchases - summary.statement_purchases
        print(
            f"{summary.source_file}: count={summary.transaction_count}, "
            f"statement_purchases={summary.statement_purchases}, "
            f"parsed_purchases={summary.parsed_purchases}, diff={diff}"
        )


if __name__ == "__main__":
    main()
